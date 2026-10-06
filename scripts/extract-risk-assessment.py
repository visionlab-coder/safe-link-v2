"""Read-only XLSX extraction. Print JSON; never edit the supplied workbook."""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
sheet = openpyxl.load_workbook(source, data_only=True).active
items = []
for row in range(11, sheet.max_row + 1):
    def cell(column):
        current = sheet.cell(row, column)
        if current.value is None:
            for merged in sheet.merged_cells.ranges:
                if current.coordinate in merged:
                    return sheet.cell(merged.min_row, merged.min_col).value
        return current.value
    hazard = cell(3)
    if not hazard:
        continue
    values = [cell(c) for c in range(1, 10)]
    assert all(values[c] is not None for c in range(8)), (row, values)
    assert all(isinstance(values[c], (int, float)) for c in (4, 5, 6)), row
    items.append(dict(
        id=f"seowon-initial-row-{row}", category=str(values[0]).strip(), subcategory=str(values[1]).strip(),
        hazard_description=str(hazard).strip(), accident_type=str(values[3]).strip(),
        frequency=int(values[4]), severity=int(values[5]), risk_level=int(values[6]),
        preventive_measure=str(values[7]).strip(), is_critical=values[8] == "●",
        source_id="seowon-initial-20261006", source_row=row,
        source_critical_mark="" if values[8] is None else str(values[8]),
    ))
print(json.dumps(dict(source=dict(id="seowon-initial-20261006", name=source.name,
    sheet=sheet.title, range="A11:I461", sha256=hashlib.sha256(source.read_bytes()).hexdigest()),
    items=items), ensure_ascii=False, indent=2))
